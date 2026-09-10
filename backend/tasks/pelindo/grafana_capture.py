"""
tasks/pelindo/grafana_capture.py
-- Weekly Capture Grafana HO --
Migrated dari: weekly_capture_grafana_HO.py

Setup:
  cp ~/template_grafana.pptx ./shared/template_grafana.pptx
"""
import os, time
from datetime import datetime
from zoneinfo import ZoneInfo
from io import BytesIO
from urllib.parse import urlparse, parse_qs

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.common.by import By

GRAFANA_USER  = os.getenv("GRAFANA_USER", "icscompute")
GRAFANA_PASS  = os.getenv("GRAFANA_PASS", "icsW4tch!")
TEMPLATE_FILE = "/app/shared/template_grafana.pptx"

URL_LIST = [
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-srvlpmy01&var-instance=10.174.3.60:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-srvlppajak01&var-instance=10.174.3.62:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-srvlppeproc01&var-instance=10.174.3.3:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-srvlpprima01&var-instance=10.174.3.53:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-srvlprkm01&var-instance=10.174.3.4:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=gcp-srvwpbios01&var-instance=10.174.3.6:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=gcp-srvwpgatewaypowerbi01&var-instance=10.174.3.5:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=gcp-srvwpsimtax&var-instance=10.174.3.61:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-srvlpcentra01&var-instance=10.174.3.199:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=gcp-svrlptanos01&var-instance=10.174.3.204:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/ferdmagigyzggb/k8s-back-office?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-back-office&var-cluster_name=back-office-new&var-namespace=&var-container_name=&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/ferdmagigyzggb/k8s-back-office?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-back-office&var-cluster_name=back-office-2&var-namespace=&var-container_name=&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/df07rysvmjx1cc/vm-linux?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-job=node-exporter&var-nodename=ematerai&var-instance=10.174.2.6:9100&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=ibs-api-1&var-instance=10.174.2.40:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=ibs-api-2&var-instance=10.174.2.42:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=integration&var-instance=10.174.2.7:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/advadwnf34b34/vm-payment?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-payment&var-instance_name=pelindo-pay&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/bey7rgxvk9b0ga/vm-windows?orgId=1&var-prome_ds=eey4rwzgb9ibke&var-nodename=payment-bank-01&var-instance=10.174.2.50:9182&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/87v4vm98cfrvr2v3/k8s-payment?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-payment&var-cluster_name=k8s-payment&var-container_name=&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/advadwnf34b34/vm-payment?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-payment&var-instance_name=All&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/ceu58oumv3qiod/k8s-sap?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-sap&var-cluster_name=sap-peic&var-namespace=&var-container_name=&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/aeu5916cfad4wb/sql-sap?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-sap&var-db_name=svreicprdsql&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/eeu5c3dxxotfkc/redis-sap?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-sap&var-instance_name=All&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/den1cqwgkwa9sa/vm-qa?orgId=1&var-datasource=cecpye14uwxkwc&var-project=pelindo-qa&var-instance_name=All&var-alignmentPeriod=grafana-auto",
    "https://monitoring.ilcs.co.id/d/eexfbo85ofwu8e/interconnect?orgId=1",
    "https://monitoring.ilcs.co.id/d/cey479vts11xce/gcp-load-balancer?orgId=1",
    "https://monitoring.ilcs.co.id/d/cefogwvbsp2bkc/cloud-armor-all?orgId=1&var-cloud_armor_name=All",
    "https://monitoring.ilcs.co.id/d/be954464vl2pseasc/vm-apps?orgId=1&var-instance_name=oci-srvlptravel01&var-interval=1m&var-compartment=pelindo%20%3E%20App-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=dlake-repo&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=dlake-edge&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=dlake-wkr&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=dlake-mstr&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=dlake-util&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=strm-kfk&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=strm-nf&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=strm-mstr&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=strm-util&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=freeipa&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/be954464vl2pse/vm-cloudera?orgId=1&var-instance_name=dns-serve&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/fegsg889m1vk0d/vm-phinnisi?orgId=1&var-instance_name=All&var-compartment=ilcspsdcloudaccount%20%3E%20prod-phinnisi&var-interval=1m&var-nodepool1=oke-c5r74d6mfva-ndgzllb7uca-s2dlrd4cwu&var-nodepool2=oke-c5r74d6mfva-nlflfmfum5q-s2dlrd4cwua",
    "https://monitoring.ilcs.co.id/d/deiv4bn1a5w5ce/db-phinnisi?orgId=1&var-db_name=All&var-compartment=ilcspsdcloudaccount%20%3E%20prod-phinnisi&var-interval=1m",
    "https://monitoring.ilcs.co.id/d/y3gun4vh5ycefng78/db-oracle-db-pelindo?orgId=1&var-db_name=madra&var-interval=1m&var-compartment=pelindo%20%3E%20DB-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/y3gun4vh5ycefng78/db-oracle-db-pelindo?orgId=1&var-db_name=pancala&var-interval=1m&var-compartment=pelindo%20%3E%20DB-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/y3gun4vh5ycefng78/db-oracle-db-pelindo?orgId=1&var-db_name=hastina&var-interval=1m&var-compartment=pelindo%20%3E%20DB-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/y3gun4vh5ycefng78/db-oracle-db-pelindo?orgId=1&var-db_name=ayodya&var-interval=1m&var-compartment=pelindo%20%3E%20DB-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/y3gun4vh5ycefng78/db-oracle-db-pelindo?orgId=1&var-db_name=alengka&var-interval=1m&var-compartment=pelindo%20%3E%20DB-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/y3gun4vh5yceftm6u/db-mysql-db-pelindo?orgId=1&var-db_name=All&var-interval=1m&var-compartment=pelindo%20%3E%20DB-Pelindo-HO",
    "https://monitoring.ilcs.co.id/d/cejrks918xtkwb/db-sigap-postgres?orgId=1&var-dbname=sigap-prod&var-Interval=1m",
    "https://monitoring.ilcs.co.id/d/cezp5v5ej4f0gc/centralized-kubenertes?orgId=1&var-interval=30m&var-cluster=huwai-portaverse&var-namespace=All&var-container=All",
    "https://monitoring.ilcs.co.id/d/cdyyfytqtixogfdsfdsf/centralized-rds?orgId=1&var-EPS=0&var-hostIP=10.95.200.227&var-hostname=dbportaverse&var-total=4",
    "https://monitoring.ilcs.co.id/d/portefaix_node_fleet_overviewdsd/centralized-vm?orgId=1&refresh=30s&var-DS_Metrics=eey51w20ssn40e&var-job=All&var-vm=simon-stid&var-rate_interval=5m",
    "https://monitoring.ilcs.co.id/d/cdyyfytqtixogfdsfdsf/centralized-rds?orgId=1&var-EPS=0&var-hostIP=10.95.200.49&var-hostname=db-stid-simon&var-total=4",
    "https://monitoring.ilcs.co.id/d/portefaix_node_fleet_overviewdsd/centralized-vm?orgId=1&refresh=30s&var-DS_Metrics=eey51w20ssn40e&var-job=All&var-vm=ims&var-rate_interval=5m",
    "https://monitoring.ilcs.co.id/d/cdyyfytqtixogfdsfdsf/centralized-rds?orgId=1&var-EPS=0&var-hostIP=10.95.200.11&var-hostname=db-apps&var-total=4",
]

SLIDE_MAPPING = {
    0:3,1:5,2:7,3:9,4:11,5:13,6:15,7:17,8:19,9:21,
    10:23,11:25,12:28,13:30,14:31,15:33,16:35,17:37,18:39,
    19:42,20:44,21:46,22:48,23:51,24:53,25:54,26:55,27:58,
    28:62,29:63,30:64,31:65,32:66,33:67,34:68,35:69,36:70,
    37:71,38:72,39:74,40:76,41:78,42:80,43:82,44:84,45:86,
    46:88,47:90,48:93,49:95,50:98,51:100,52:102,53:104,
}


def _get_hostname(url):
    params = parse_qs(urlparse(url).query)
    for key in ["var-nodename", "var-instance_name", "var-instance", "var-vm"]:
        if key in params:
            val = params[key][0]
            if ":" in val: val = val.split(":")[0]
            if val.lower() != "all": return val
    return "dashboard"


def run(log_path=None, args=None):
    def log(msg):
        if log_path:
            with open(log_path, "a") as f:
                f.write(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}\n")
        print(msg)

    args       = args or {}
    start_date = args.get("start_date", datetime.now().strftime("%Y-%m-%d"))
    start_time = args.get("start_time", "00:00")
    end_date   = args.get("end_date",   datetime.now().strftime("%Y-%m-%d"))
    end_time   = args.get("end_time",   "23:59")

    def to_ts(d, t):
        dt = datetime.strptime(f"{d} {t}", "%Y-%m-%d %H:%M").replace(tzinfo=ZoneInfo("Asia/Jakarta"))
        return int(dt.timestamp() * 1000)

    FROM_TS = to_ts(start_date, start_time)
    TO_TS   = to_ts(end_date, end_time)

    if not os.path.exists(TEMPLATE_FILE):
        return f"ERROR: template_grafana.pptx tidak ada. Jalankan: cp ~/template_grafana.pptx ./shared/"

    from chrome_helper import get_driver
    driver = get_driver(1920, 1080)
    wait   = WebDriverWait(driver, 60)
    captured = []

    try:
        log("Login Grafana...")
        driver.get("https://monitoring.ilcs.co.id/login")
        wait.until(lambda d: d.find_element(By.NAME, "user"))
        driver.find_element(By.NAME, "user").send_keys(GRAFANA_USER)
        driver.find_element(By.NAME, "password").send_keys(GRAFANA_PASS)
        driver.find_element(By.CSS_SELECTOR, "button[type='submit']").click()
        time.sleep(8)
        log("Login OK")

        for idx, base_url in enumerate(URL_LIST):
            sep = "&" if "?" in base_url else "?"
            url  = f"{base_url}{sep}from={FROM_TS}&to={TO_TS}&timezone=Asia%2FJakarta"
            name = _get_hostname(base_url)
            log(f"[{idx+1}/{len(URL_LIST)}] {name}")

            driver.get(url)
            wait.until(lambda d: d.execute_script("return document.readyState") == "complete")
            time.sleep(10)
            driver.execute_script("window.scrollTo(0,0)")
            time.sleep(2)
            driver.execute_script("document.body.style.zoom='90%'")
            time.sleep(2)

            png   = driver.get_screenshot_as_png()
            image = Image.open(BytesIO(png))
            w, h  = image.size
            image = image.crop((0, 0, w, int(h * 0.87)))
            captured.append((name, image))

    finally:
        driver.quit()

    log(f"Total captured: {len(captured)} — Generating PPTX...")

    output_dir = "/app/shared/reports"
    os.makedirs(output_dir, exist_ok=True)
    ppt_name = f"{output_dir}/pelindo_grafana_{start_date}_to_{end_date}.pptx"

    if os.path.exists(ppt_name):
        os.remove(ppt_name)

    prs = Presentation(TEMPLATE_FILE)

    for idx, (name, image) in enumerate(captured):
        if idx not in SLIDE_MAPPING:
            continue
        slide = prs.slides[SLIDE_MAPPING[idx]]
        if slide.shapes.title:
            slide.shapes.title.text = f"{name} ({start_date} to {end_date})"
        stream = BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
        slide.shapes.add_picture(stream, Inches(0.3), Inches(1.3), width=Inches(12.8))

    prs.save(ppt_name)
    log(f"Saved: {ppt_name}")

    # Hapus file lama, sisakan 2 terbaru
    try:
        all_files = sorted(
            [f for f in os.listdir(output_dir) if f.startswith("pelindo_grafana_") and f.endswith(".pptx")],
            reverse=True
        )
        for old in all_files[2:]:
            if old != os.path.basename(ppt_name):
                os.remove(os.path.join(output_dir, old))
                log(f"Deleted old file: {old}")
    except Exception as e:
        log(f"Cleanup warning: {e}")

    return f"DONE: {os.path.basename(ppt_name)}"
